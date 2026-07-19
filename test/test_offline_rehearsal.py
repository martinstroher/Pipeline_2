"""End-to-end regression for the strictly offline evaluation rehearsal.

Run:
    python test/test_offline_rehearsal.py
"""

from __future__ import annotations

import json
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parent.parent
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from src.evaluation.offline_rehearsal import (  # noqa: E402
    PROVENANCE,
    _genericise_definition,
    _sha256_file,
    run_offline_rehearsal,
)


def _run(path: Path) -> dict:
    return run_offline_rehearsal(
        output_dir=str(path),
        bootstrap_iterations=100,
        overwrite=False,
    )


def main() -> int:
    generic = _genericise_definition(
        "Santos Basin",
        "Santos Basin is a Brazilian Pre-Salt basin. It has a second sentence.",
    )
    assert generic == "Santos Basin is a petroleum-system basin."
    assert _genericise_definition("Term", "No final punctuation") == "No final punctuation."

    with tempfile.TemporaryDirectory() as directory:
        root = Path(directory)
        first = Path(_run(root / "first_rehearsal")["output_dir"])
        second = Path(_run(root / "second_rehearsal")["output_dir"])

        deterministic_files = [
            "nld_B.csv",
            "nld_C.csv",
            "nld_D.csv",
            "cat_B.csv",
            "cat_C.csv",
            "cat_D.csv",
            "ablation_merged.csv",
            "analysis/layer1/sensitivity_flags.csv",
        ]
        for relative_path in deterministic_files:
            assert _sha256_file(first / relative_path) == _sha256_file(second / relative_path), relative_path

        first_analysis = json.loads(
            (first / "analysis/layer2/layer2_results.json").read_text(encoding="utf-8")
        )
        second_analysis = json.loads(
            (second / "analysis/layer2/layer2_results.json").read_text(encoding="utf-8")
        )
        assert first_analysis["analysis_design"]["source_hashes_validated"] is True
        assert second_analysis["analysis_design"]["source_hashes_validated"] is True
        first_analysis["analysis_design"].pop("source_manifest")
        second_analysis["analysis_design"].pop("source_manifest")
        assert first_analysis == second_analysis

        manifest = json.loads(
            (first / "OFFLINE_REHEARSAL_MANIFEST.json").read_text(encoding="utf-8")
        )
        assert manifest["artifact_kind"] == PROVENANCE
        assert manifest["azure_calls"] == 0
        assert manifest["human_experts"] == 0
        assert manifest["scientific_inference_permitted"] is False
        assert manifest["generated_file_sha256"]["cat_D.csv"] == _sha256_file(first / "cat_D.csv")

        workbooks = sorted((first / "expert_workbooks").glob("*.xlsx"))
        assert len(workbooks) == 3
        assert (first / "private" / "blinding_key_42.csv").exists()
        assert not list((first / "expert_workbooks").glob("*key*"))

        workbook = load_workbook(workbooks[0])
        assert workbook.sheetnames[0] == "REHEARSAL_ONLY"
        representation = workbook["Representation"]
        relevance_column = next(
            cell.column_letter
            for cell in representation[1]
            if cell.value == "Relevance (1-5)"
        )
        validations = list(representation.data_validations.dataValidation)
        assert any(
            relevance_column in str(validation.sqref) and validation.allow_blank is False
            for validation in validations
        )
        assert len(representation.conditional_formatting) > 0

    print("=== OFFLINE REHEARSAL TEST PASSED ===")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())