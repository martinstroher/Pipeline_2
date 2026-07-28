"""End-to-end regression for the strictly offline evaluation rehearsal.

Run:
    python evaluation_study/test/test_offline_rehearsal.py
"""

from __future__ import annotations

import json
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[2]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from evaluation_study.offline_rehearsal import (  # noqa: E402
    PROVENANCE,
    _genericise_definition,
    _sha256_file,
    run_offline_rehearsal,
)
from evaluation_study.expert_eval_workbook import (  # noqa: E402
    DATA_HEADER_ROW,
    DATA_START_ROW,
    visible_column_name,
    visible_sheet_name,
)


def _run(path: Path) -> dict:
    return run_offline_rehearsal(
        output_dir=str(path),
        bootstrap_iterations=100,
        overwrite=False,
    )


def main() -> int:
    generic = _genericise_definition(
        "Santos Basin",
        "Santos Basin is a Brazilian Pre-Salt basin. It has a second sentence.",
    )
    assert generic == "Santos Basin is a petroleum-system basin."
    assert _genericise_definition("Term", "No final punctuation") == "No final punctuation."

    with tempfile.TemporaryDirectory() as directory:
        root = Path(directory)
        first = Path(_run(root / "first_rehearsal")["output_dir"])
        second = Path(_run(root / "second_rehearsal")["output_dir"])

        deterministic_files = [
            "nld_B.csv",
            "nld_C.csv",
            "nld_D.csv",
            "cat_B.csv",
            "cat_C.csv",
            "cat_D.csv",
            "ablation_merged.csv",
            "analysis/layer1/sensitivity_flags.csv",
            "analysis/layer2/discordant_category_contrasts.csv",
            "analysis/layer2/cross_layer_spearman.csv",
            "analysis/layer2/item_consensus.csv",
            "analysis/layer2/consensus_summary.csv",
        ]
        for relative_path in deterministic_files:
            assert _sha256_file(first / relative_path) == _sha256_file(second / relative_path), relative_path

        first_analysis = json.loads(
            (first / "analysis/layer2/layer2_results.json").read_text(encoding="utf-8")
        )
        second_analysis = json.loads(
            (second / "analysis/layer2/layer2_results.json").read_text(encoding="utf-8")
        )
        assert first_analysis["analysis_design"]["source_hashes_validated"] is True
        assert second_analysis["analysis_design"]["source_hashes_validated"] is True
        assert first_analysis["analysis_design"]["response_handling"]["ties"] == "retained and counted explicitly"
        assert "meaning_preservation" not in first_analysis["final_ontology"]
        assert "core_appropriateness" not in first_analysis["final_ontology"]
        assert first_analysis["exploratory_cross_layer"]["status"] == "exploratory"
        assert first_analysis["consensus"]["item_table"] == "item_consensus.csv"
        assert first_analysis["completion_time"]["overall_mean_minutes"] > 0
        first_analysis["analysis_design"].pop("source_manifest")
        second_analysis["analysis_design"].pop("source_manifest")
        assert first_analysis == second_analysis

        manifest = json.loads(
            (first / "OFFLINE_REHEARSAL_MANIFEST.json").read_text(encoding="utf-8")
        )
        assert manifest["artifact_kind"] == PROVENANCE
        assert manifest["azure_calls"] == 0
        assert manifest["human_experts"] == 0
        assert manifest["scientific_inference_permitted"] is False
        assert manifest["generated_file_sha256"]["cat_D.csv"] == _sha256_file(first / "cat_D.csv")

        workbooks = sorted((first / "expert_workbooks").glob("*.xlsx"))
        assert len(workbooks) == 3
        assert (first / "private" / "blinding_key_42.csv").exists()
        assert not list((first / "expert_workbooks").glob("*key*"))

        workbook = load_workbook(workbooks[0])
        assert workbook.sheetnames[0] == "REHEARSAL_ONLY"
        assert "Practice" not in workbook.sheetnames
        representation = workbook["Representation"]
        relevance_column = next(
            cell.column_letter
            for cell in representation[DATA_HEADER_ROW]
            if cell.value == visible_column_name(
                "Representation", "Relevance (1-5/Unsure)"
            )
        )
        validations = list(representation.data_validations.dataValidation)
        assert any(
            relevance_column in str(validation.sqref) and validation.allow_blank is False
            for validation in validations
        )
        assert len(representation.conditional_formatting) > 0
        assert representation.column_dimensions["A"].hidden is True

        response_sheets = (
            "Representation",
            "Category_Correct",
            "Taxonomy",
            "Defined_Classes",
            "Relations",
            "Individuals",
        )
        expert_facing_sheets = (
            "Category_Guide",
            *response_sheets,
            "Timing",
        )
        for sheet_name in expert_facing_sheets:
            sheet = workbook[visible_sheet_name(sheet_name)]
            assert str(sheet["A1"].value).strip()
            assert str(sheet["A2"].value).startswith("YOUR TASK\n")
            assert str(sheet["A3"].value).startswith("HOW TO ANSWER\n")
            if sheet_name == "Category_Guide":
                assert "nothing to fill in" in str(sheet["A4"].value)
            elif sheet_name == "Timing":
                assert "Enter the minutes" in str(sheet["A4"].value)
            else:
                assert "Yellow cells" in str(sheet["A4"].value)
            assert sheet.row_dimensions[1].height >= 30
            assert sheet.row_dimensions[2].height >= 45
            assert sheet.row_dimensions[3].height >= 55

        row_ids_by_workbook = []
        for path in workbooks:
            candidate = load_workbook(path, read_only=True)
            row_ids_by_workbook.append({
                sheet_name: {
                    str(row[0])
                    for row in candidate[visible_sheet_name(sheet_name)].iter_rows(
                        min_row=DATA_START_ROW,
                        values_only=True,
                    )
                }
                for sheet_name in response_sheets
            })
        assert row_ids_by_workbook[0] == row_ids_by_workbook[1]
        assert row_ids_by_workbook[1] == row_ids_by_workbook[2]

        workbook_manifest = json.loads(
            (first / "expert_evaluation_manifest.json").read_text(encoding="utf-8")
        )
        assert "assignment_design" not in workbook_manifest
        optional_review_path = Path(
            workbook_manifest["optional_model_changes_review"]
        )
        assert optional_review_path.exists()
        optional_review = load_workbook(optional_review_path)
        assert optional_review.sheetnames == [
            "Instructions",
            visible_sheet_name("Meaning_Preservation"),
        ]
        assert optional_review["Instructions"]["A2"].value == (
            "OPTIONAL SPECIALIST REVIEW"
        )
        assert optional_review["Instructions"]["B2"].value == (
            "This separate file is for a geologist or ontologist who is comfortable "
            "reviewing how terms were removed or rewritten. It is not part of the "
            "main three-expert study."
        )
        optional_sheet = optional_review[
            visible_sheet_name("Meaning_Preservation")
        ]
        assert optional_sheet.max_row - DATA_HEADER_ROW == 40
        optional_headers = {
            cell.value: cell.column for cell in optional_sheet[DATA_HEADER_ROW]
        }
        for header in (
            visible_column_name(
                "Meaning_Preservation",
                "Meaning_Preserved (Fully/Mostly/No/Unsure)",
            ),
            visible_column_name(
                "Meaning_Preservation",
                "Appropriate_for_Lean_Core (Yes/With concern/No/Unsure)",
            ),
        ):
            assert all(
                optional_sheet.cell(row, optional_headers[header]).value in (None, "")
                for row in range(DATA_START_ROW, optional_sheet.max_row + 1)
            )

        expected_absent = {
            "Category_Correct": {"Suggested_Category"},
            "Taxonomy": {"Suggested_Parent", "Question", "Keep_in_Core (Yes/No/Unsure)"},
            "Defined_Classes": {"Suggested_Change", "Characteristic_General (Yes/No/Unsure)", "Feature_Is_Defining_in_PreSalt (Yes/No/Unsure)"},
            "Relations": {"Statement_Correct (Yes/Partial/No/Unsure)", "Generally_True (Yes/No/Unsure)"},
            "Individuals": {"Suggested_Type", "Rationale"},
        }
        for sheet_name, absent_headers in expected_absent.items():
            headers = {
                cell.value
                for cell in workbook[visible_sheet_name(sheet_name)][DATA_HEADER_ROW]
            }
            assert not headers & absent_headers, (sheet_name, headers & absent_headers)

        assert visible_sheet_name("Category_Guide") in workbook.sheetnames
        assert visible_column_name(
            "Taxonomy", "Useful_PreSalt_Distinction (Yes/No/Unsure)"
        ) in {
            cell.value for cell in workbook["Taxonomy"][DATA_HEADER_ROW]
        }
        assert visible_column_name(
            "Defined_Classes",
            "Definition_Verdict (Correct/Partly correct/Incorrect/Unsure)",
        ) in {
            cell.value
            for cell in workbook[visible_sheet_name("Defined_Classes")][DATA_HEADER_ROW]
        }
        assert visible_column_name("Relations", "Relation_Verdict") in {
            cell.value for cell in workbook["Relations"][DATA_HEADER_ROW]
        }
        reference_header = visible_column_name(
            "Category_Correct", "Reference_Definition"
        )
        category_sheet = workbook[visible_sheet_name("Category_Correct")]
        category_headers = {
            cell.value: cell.column
            for cell in category_sheet[DATA_HEADER_ROW]
        }
        assert reference_header in category_headers
        reference_values = [
            category_sheet.cell(row, category_headers[reference_header]).value
            for row in range(DATA_START_ROW, category_sheet.max_row + 1)
        ]
        assert all(value not in (None, "") for value in reference_values)
        for sheet_name in response_sheets[2:]:
            assert "Term_Context (only when needed)" not in {
                cell.value
                for cell in workbook[visible_sheet_name(sheet_name)][DATA_HEADER_ROW]
            }

        relation_sheet = workbook["Relations"]
        relation_headers = {
            cell.value: cell.column
            for cell in relation_sheet[DATA_HEADER_ROW]
        }
        relation_statement_header = visible_column_name(
            "Relations", "Relation_Statement"
        )
        relation_statements = {
            str(relation_sheet.cell(row, relation_headers[relation_statement_header]).value)
            for row in range(DATA_START_ROW, relation_sheet.max_row + 1)
        }
        assert any(value.startswith("Generally, ") for value in relation_statements)
        assert any(
            value.startswith("In some reported Pre-Salt contexts, ")
            for value in relation_statements
        )
        assert any(
            value.startswith("For this named entity, ")
            for value in relation_statements
        )
        assert visible_sheet_name("Meaning_Preservation") not in workbook.sheetnames

    print("=== OFFLINE REHEARSAL TEST PASSED ===")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())