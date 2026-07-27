"""Strictly offline end-to-end rehearsal of the thesis evaluation workflow.

This module does not reproduce GPT-5.4 or human judgments. It creates clearly
marked, deterministic surrogate B/C/D outputs and mock expert responses solely
to exercise schemas, manifests, statistics, workbook blinding, unblinding, and
reporting before paid model execution or expert review.

No function in this module imports or calls the Azure LLM client.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import re
import shutil
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

import evaluation_study.layer1_analysis as layer1
from evaluation_study.expert_eval_analysis import run_modular_analysis
from evaluation_study.expert_eval_generator import generate_expert_evaluation
from evaluation_study.expert_eval_workbook import (
    DATA_HEADER_ROW,
    DATA_START_ROW,
    StudyInputs,
    build_final_fates,
)
from evaluation_study.paths import (
    APPROVED_ONTOLOGY_DIR,
    FILTERED_TERMS,
    FROZEN_A_CATEGORIES,
    FROZEN_A_NLD,
    REHEARSAL_OUTPUT,
)
from src.utils.csv_io import read_csv, write_csv
from src.utils.ontology_config import get_config


PROVENANCE = "SYNTHETIC_OFFLINE_REHEARSAL"
WARNING = (
    "SYNTHETIC OFFLINE REHEARSAL ONLY. No Azure model and no human expert "
    "produced these results. Do not use them as thesis evidence."
)
DEFAULT_OUTPUT_DIR = str(REHEARSAL_OUTPUT)
DEFAULT_ONTOLOGY_DIR = str(APPROVED_ONTOLOGY_DIR)
DEFAULT_TERMS_PATH = str(FILTERED_TERMS)
DEFAULT_A_NLD_PATH = str(FROZEN_A_NLD)
DEFAULT_A_CATEGORY_PATH = str(FROZEN_A_CATEGORIES)
EXPECTED_TERM_COUNT = 407
SEED = 42
CONDITIONS = ("A", "B", "C", "D")

CHANGE_RATES = {"B": 0.14, "C": 0.27, "D": 0.11}
SAME_TIER_PROBABILITY = {"B": 0.75, "C": 0.55, "D": 0.65}
NOT_CLASSIFIED_PROBABILITY = {"B": 0.03, "C": 0.08, "D": 0.04}
FATE_MULTIPLIERS = {
    "FINAL_CLASS": 0.75,
    "FINAL_INDIVIDUAL": 0.85,
    "DEMOTED": 1.35,
    "CRITIC_EXCLUDED": 1.25,
    "CQ_EXCLUDED": 1.45,
}


def _normalise(value: object) -> str:
    return str(value).strip().casefold()


def _sha256_file(path: str | Path) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _sha256_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def _unit(seed: int, *parts: object) -> float:
    payload = "|".join([str(seed), *[str(part) for part in parts]])
    integer = int(hashlib.sha256(payload.encode("utf-8")).hexdigest()[:16], 16)
    return integer / float(16**16 - 1)


def _pick(options: list[str], seed: int, *parts: object) -> str:
    if not options:
        raise ValueError("Cannot choose from an empty option list")
    index = min(int(_unit(seed, *parts) * len(options)), len(options) - 1)
    return options[index]


def _prepare_output_dir(path: str | Path, overwrite: bool) -> Path:
    output_dir = Path(path).resolve()
    if "rehearsal" not in output_dir.name.casefold():
        raise ValueError(
            f"Offline rehearsal output directory must contain 'rehearsal': {output_dir}"
        )
    if output_dir.exists():
        if not overwrite:
            raise FileExistsError(
                f"Rehearsal directory already exists: {output_dir}. Use --overwrite."
            )
        print(f"Removing existing rehearsal directory: {output_dir}")
        shutil.rmtree(output_dir)
    output_dir.mkdir(parents=True)
    return output_dir


def _require_columns(frame: pd.DataFrame, columns: set[str], label: str) -> None:
    missing = columns - set(frame.columns)
    if missing:
        raise ValueError(f"{label} missing columns: {sorted(missing)}")


def _load_source_inputs(
    terms_path: str,
    a_nld_path: str,
    a_category_path: str,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    terms = read_csv(terms_path)
    nld_a = read_csv(a_nld_path)
    cat_a = read_csv(a_category_path)
    _require_columns(terms, {"Readable_Term", "Frequency"}, "Filtered terms")
    _require_columns(nld_a, {"Term", "NLD", "Context_Used", "Context"}, "Frozen A NLD")
    _require_columns(cat_a, {"Term", "Category", "Reasoning", "NLD"}, "Frozen A categories")

    if len(terms) != EXPECTED_TERM_COUNT:
        raise ValueError(f"Expected {EXPECTED_TERM_COUNT} terms; found {len(terms)}")
    for label, frame, column in (
        ("Filtered terms", terms, "Readable_Term"),
        ("Frozen A NLD", nld_a, "Term"),
        ("Frozen A categories", cat_a, "Term"),
    ):
        if frame[column].astype(str).duplicated().any():
            raise ValueError(f"{label} contains duplicate terms")
    term_set = set(terms["Readable_Term"].astype(str))
    if set(nld_a["Term"].astype(str)) != term_set:
        raise ValueError("Frozen A NLD term set does not match filtered terms")
    if set(cat_a["Term"].astype(str)) != term_set:
        raise ValueError("Frozen A category term set does not match filtered terms")
    if nld_a["Context"].isna().any() or nld_a["Context"].astype(str).str.strip().eq("").any():
        missing = nld_a.loc[
            nld_a["Context"].isna() | nld_a["Context"].astype(str).str.strip().eq(""),
            "Term",
        ].tolist()
        raise ValueError(f"Frozen A has missing contexts required by D: {missing[:10]}")
    errors = cat_a["Category"].fillna("").astype(str).str.startswith("ERROR")
    if errors.any():
        raise ValueError(f"Frozen A contains {int(errors.sum())} category error rows")
    return terms, nld_a, cat_a


def _genericise_definition(term: str, definition: object) -> str:
    """Create a short B proxy from A's first sentence.

    This intentionally leaks A and is therefore suitable only for process
    rehearsal. The limitation is recorded in every manifest and report.
    """
    text = re.sub(r"\s+", " ", str(definition)).strip()
    if not text:
        return f"{term} is a geological concept."
    sentence_match = re.match(r"(.+?[.!?])(?:\s|$)", text)
    sentence = sentence_match.group(1) if sentence_match else text
    sentence = re.sub(
        r"\b(?:Brazilian|South Atlantic)\s+Pre[- ]Salt\b",
        "petroleum-system",
        sentence,
        flags=re.IGNORECASE,
    )
    sentence = re.sub(r"\s+", " ", sentence).strip()
    if not sentence.endswith((".", "!", "?")):
        sentence += "."
    return sentence


def _build_nld_surrogates(
    output_dir: Path,
    nld_a: pd.DataFrame,
    a_nld_path: str,
) -> dict[str, pd.DataFrame]:
    shutil.copy2(a_nld_path, output_dir / "nld_A.csv")

    nld_b = pd.DataFrame({
        "Term": nld_a["Term"],
        "NLD": [
            _genericise_definition(term, definition)
            for term, definition in zip(nld_a["Term"], nld_a["NLD"])
        ],
        "Context_Used": False,
        "Context": "",
        "Synthetic_Rehearsal": True,
        "Proxy_Method": "first_sentence_of_frozen_A",
    })
    nld_c = pd.DataFrame({
        "Term": nld_a["Term"],
        "NLD": "",
        "Context_Used": False,
        "Context": "",
        "Synthetic_Rehearsal": True,
        "Proxy_Method": "empty_nld_term_only",
    })
    nld_d = pd.DataFrame({
        "Term": nld_a["Term"],
        "NLD": nld_a["Context"],
        "Context_Used": nld_a["Context_Used"],
        "Context": nld_a["Context"],
        "Synthetic_Rehearsal": True,
        "Proxy_Method": "exact_frozen_A_context",
    })
    for condition, frame in (("B", nld_b), ("C", nld_c), ("D", nld_d)):
        write_csv(frame, output_dir / f"nld_{condition}.csv")
    return {"A": nld_a, "B": nld_b, "C": nld_c, "D": nld_d}


def _fate_inputs(
    terms: pd.DataFrame,
    ontology_dir: str,
) -> StudyInputs:
    root = Path(ontology_dir)
    return StudyInputs(
        terms=terms,
        nld={},
        categories={},
        refined_categories=read_csv(root / "classify_categories.csv"),
        taxonomy=read_csv(root / "validate_taxonomy.csv"),
        defined_classes=read_csv(root / "validate_defined_classes.csv"),
        relations=read_csv(root / "validate_relations.csv"),
        individuals=read_csv(root / "validate_instances.csv"),
        demotions=read_csv(root / "validate_demotions.csv"),
        class_fates=read_csv(root / "validate_class_fates.csv"),
    )


def _reference_categories(
    cat_a: pd.DataFrame,
    ontology_dir: str,
) -> dict[str, str]:
    reference = {
        _normalise(term): str(category)
        for term, category in zip(cat_a["Term"], cat_a["Category"])
    }
    root = Path(ontology_dir)
    sources = (
        (root / "validate_taxonomy.csv", "Term", "Category"),
        (root / "validate_instances.csv", "Term", "Original_Category"),
        (root / "validate_demotions.csv", "Term", "Original_Category"),
        (root / "validate_class_fates.csv", "term", "category"),
    )
    for path, term_column, category_column in sources:
        frame = read_csv(path)
        if term_column not in frame.columns or category_column not in frame.columns:
            continue
        for term, category in zip(frame[term_column], frame[category_column]):
            if pd.notna(category):
                reference[_normalise(term)] = str(category)
    return reference


def _category_metadata() -> tuple[set[str], dict[str, str], dict[str, frozenset[str]]]:
    config = get_config()
    valid = {"NOT_CLASSIFIED"}
    tiers = {"NOT_CLASSIFIED": "NOT_CLASSIFIED"}
    for ontology_key in config.waterfall_ontologies():
        tier = config.ontologies[ontology_key].eval_tier.upper()
        for category in config.categories_for(ontology_key):
            valid.add(category)
            tiers[category] = tier
    return valid, tiers, config.category_to_metatypes()


def _metatype_overlap(
    first: str,
    second: str,
    metatypes: dict[str, frozenset[str]],
) -> float:
    first_types = set(metatypes.get(first, ()))
    second_types = set(metatypes.get(second, ()))
    union = first_types | second_types
    return len(first_types & second_types) / len(union) if union else 0.0


def _alternative_category(
    condition: str,
    term: str,
    current: str,
    valid: set[str],
    tiers: dict[str, str],
    metatypes: dict[str, frozenset[str]],
    seed: int,
) -> str:
    if (
        current != "NOT_CLASSIFIED"
        and _unit(seed, condition, term, "not-classified")
        < NOT_CLASSIFIED_PROBABILITY[condition]
    ):
        return "NOT_CLASSIFIED"

    categories = sorted(valid - {current, "NOT_CLASSIFIED"})
    same_tier = (
        _unit(seed, condition, term, "same-tier")
        < SAME_TIER_PROBABILITY[condition]
    )
    if same_tier:
        candidates = [category for category in categories if tiers[category] == tiers[current]]
    else:
        candidates = [category for category in categories if tiers[category] != tiers[current]]
    if not candidates:
        candidates = categories
    ranked = sorted(
        candidates,
        key=lambda category: (-_metatype_overlap(current, category, metatypes), category),
    )
    return _pick(ranked[: min(6, len(ranked))], seed, condition, term, "alternative")


def _build_category_surrogates(
    output_dir: Path,
    terms: pd.DataFrame,
    nlds: dict[str, pd.DataFrame],
    cat_a: pd.DataFrame,
    a_category_path: str,
    ontology_dir: str,
    seed: int,
) -> tuple[dict[str, pd.DataFrame], dict[str, str], dict[str, str]]:
    shutil.copy2(a_category_path, output_dir / "cat_A.csv")
    fates = build_final_fates(_fate_inputs(terms, ontology_dir))
    reference = _reference_categories(cat_a, ontology_dir)
    valid, tiers, metatypes = _category_metadata()
    invalid_a = sorted(set(cat_a["Category"].astype(str)) - valid)
    if invalid_a:
        raise ValueError(f"Frozen A contains invalid categories: {invalid_a}")

    a_lookup = cat_a.set_index("Term")
    results: dict[str, pd.DataFrame] = {"A": cat_a}
    for condition in ("B", "C", "D"):
        rows = []
        nld_lookup = nlds[condition].set_index("Term")
        for term in terms["Readable_Term"].astype(str):
            current = str(a_lookup.loc[term, "Category"])
            fate = fates[_normalise(term)]
            probability = min(0.75, CHANGE_RATES[condition] * FATE_MULTIPLIERS[fate])
            if condition == "B" and not bool(nlds["A"].set_index("Term").loc[term, "Context_Used"]):
                probability = min(probability, 0.04)
            changed = _unit(seed, condition, term, "change") < probability
            category = (
                _alternative_category(
                    condition,
                    term,
                    current,
                    valid,
                    tiers,
                    metatypes,
                    seed,
                )
                if changed
                else current
            )
            nld_row = nld_lookup.loc[term]
            rows.append({
                "Term": term,
                "Category": category,
                "Reasoning": (
                    f"[{PROVENANCE}] Deterministic process-test proxy; "
                    f"A={current}; fate={fate}; changed={changed}. NOT SCIENTIFIC DATA."
                ),
                "NLD": nld_row["NLD"],
                "Context_Used": nld_row["Context_Used"],
                "Condition": condition,
                "Synthetic_Rehearsal": True,
                "Reference_A_Category": current,
                "Pseudo_Reference_Category": reference[_normalise(term)],
                "Final_Fate": fate,
                "Perturbed": changed,
            })
        frame = pd.DataFrame(rows)
        if not set(frame["Category"]).issubset(valid):
            raise ValueError(f"Condition {condition} generated invalid categories")
        write_csv(frame, output_dir / f"cat_{condition}.csv")
        results[condition] = frame
    return results, fates, reference


def _write_merged(
    output_dir: Path,
    categories: dict[str, pd.DataFrame],
) -> pd.DataFrame:
    _, tiers, _ = _category_metadata()
    frames = []
    for condition in CONDITIONS:
        frame = categories[condition].copy()
        frame["Condition"] = condition
        frame["Tier"] = frame["Category"].map(tiers)
        frame["Synthetic_Rehearsal"] = condition != "A"
        frames.append(frame)
    merged = pd.concat(frames, ignore_index=True, sort=False)
    if len(merged) != EXPECTED_TERM_COUNT * len(CONDITIONS):
        raise ValueError(f"Merged rehearsal has {len(merged)} rows; expected 1628")
    if merged.duplicated(["Term", "Condition"]).any():
        raise ValueError("Merged rehearsal has duplicate term-condition rows")
    write_csv(merged, output_dir / "ablation_merged.csv")
    return merged


def _run_layer1(output_dir: Path) -> dict:
    previous_analysis_dir = layer1.ANALYSIS_DIR
    layer1_dir = output_dir / "analysis" / "layer1"
    layer1.ANALYSIS_DIR = str(layer1_dir)
    try:
        return layer1.run_layer1_analysis(
            merged_path=str(output_dir / "ablation_merged.csv"),
            expected_term_count=EXPECTED_TERM_COUNT,
        )
    finally:
        layer1.ANALYSIS_DIR = previous_analysis_dir


def _rating(base: float, expert: str, row_id: str, dimension: str, seed: int) -> int:
    expert_number = int(expert.rsplit("_", 1)[-1])
    expert_bias = {1: -0.10, 2: 0.10, 3: 0.0}.get(
        expert_number,
        ((expert_number % 5) - 2) * 0.04,
    )
    noise = (_unit(seed, expert, row_id, dimension) - 0.5) * 1.4
    return int(np.clip(round(base + expert_bias + noise), 1, 5))


def _verdict(
    base_level: int,
    expert: str,
    row_id: str,
    dimension: str,
    seed: int,
    partial: bool = True,
) -> str:
    expert_number = int(expert.rsplit("_", 1)[-1])
    unsure_probability = {1: 0.02, 2: 0.04, 3: 0.08}.get(
        expert_number,
        0.02 + (expert_number % 5) * 0.01,
    )
    if _unit(seed, expert, row_id, dimension, "unsure") < unsure_probability:
        return "Unsure"
    level = base_level
    shift = _unit(seed, expert, row_id, dimension, "shift")
    if shift < 0.16:
        level -= 1
    elif shift > 0.92:
        level += 1
    if partial:
        return {0: "No", 1: "Partial", 2: "Yes"}[int(np.clip(level, 0, 2))]
    return "Yes" if level >= 2 else "No"


def _relevance_base(hidden: pd.Series) -> float:
    fate_base = {
        "FINAL_CLASS": 4.3,
        "FINAL_INDIVIDUAL": 4.1,
        "DEMOTED": 3.4,
        "CRITIC_EXCLUDED": 2.8,
        "CQ_EXCLUDED": 2.3,
    }
    tier_adjustment = {
        "GEORESERVOIR": 0.25,
        "GEOCORE": 0.10,
        "BFO": -0.10,
        "NOT_CLASSIFIED": -0.25,
    }
    frequency_adjustment = {"High": 0.20, "Middle": 0.0, "Low": -0.20}
    return (
        fate_base[str(hidden["Final_Fate"])]
        + tier_adjustment.get(str(hidden["Tier_A"]), 0.0)
        + frequency_adjustment.get(str(hidden["Frequency_Band"]), 0.0)
    )


def _definition_base(
    condition: str,
    hidden: pd.Series,
    row_id: str,
    seed: int,
) -> float:
    shared_item_effect = (_unit(seed, row_id, "definition-item") - 0.5) * 1.4
    if condition == "A":
        return (
            4.05
            + shared_item_effect
            + (0.20 if str(hidden.get("Context_Used_A", "")).lower() == "true" else 0.0)
        )
    return 3.55 + shared_item_effect


def _add_rehearsal_warning(workbook, expert_id: str) -> None:
    if "REHEARSAL_ONLY" in workbook.sheetnames:
        del workbook["REHEARSAL_ONLY"]
    sheet = workbook.create_sheet("REHEARSAL_ONLY", 0)
    sheet["A1"] = "SYNTHETIC OFFLINE REHEARSAL"
    sheet["A2"] = WARNING
    sheet["A3"] = f"Workbook mock profile: {expert_id}"
    sheet["A4"] = "All ratings were generated by deterministic rules, not people."
    sheet["A1"].font = Font(bold=True, color="FFFFFF", size=18)
    sheet["A1"].fill = PatternFill("solid", fgColor="C00000")
    for cell in ("A2", "A3", "A4"):
        sheet[cell].font = Font(bold=True, color="C00000")
        sheet[cell].alignment = Alignment(wrap_text=True)
    sheet.column_dimensions["A"].width = 110


def _fill_representation_sheet(
    sheet,
    hidden: pd.DataFrame,
    expert_id: str,
    seed: int,
) -> None:
    headers = {cell.value: cell.column for cell in sheet[DATA_HEADER_ROW]}
    key = hidden.set_index("Row_ID")
    for row_index in range(DATA_START_ROW, sheet.max_row + 1):
        row_id = str(sheet.cell(row_index, headers["Row_ID"]).value)
        item = key.loc[row_id]
        relevance = _rating(_relevance_base(item), expert_id, row_id, "relevance", seed)
        quality_by_condition = {
            condition: _rating(
                _definition_base(condition, item, row_id, seed),
                expert_id,
                row_id,
                f"quality-{condition}",
                seed,
            )
            for condition in ("A", "B")
        }
        first_condition = str(item["Definition_1_Condition"])
        second_condition = str(item["Definition_2_Condition"])
        quality_first = quality_by_condition[first_condition]
        quality_second = quality_by_condition[second_condition]
        if quality_first == quality_second:
            preference = "Tie"
        elif abs(quality_first - quality_second) == 1 and _unit(
            seed, expert_id, row_id, "preference-tie"
        ) < 0.25:
            preference = "Tie"
        else:
            preference = "1" if quality_first > quality_second else "2"
        sheet.cell(row_index, headers["Relevance (1-5/Unsure)"], relevance)
        sheet.cell(row_index, headers["Quality_1 (1-5/Unsure)"], quality_first)
        sheet.cell(row_index, headers["Quality_2 (1-5/Unsure)"], quality_second)
        sheet.cell(row_index, headers["Preference (1/2/Tie/Unsure)"], preference)
        sheet.cell(row_index, headers["Notes"], f"[{PROVENANCE}; {expert_id}]")


def _fill_category_sheet(
    sheet,
    hidden: pd.DataFrame,
    expert_id: str,
    reference: dict[str, str],
    tiers: dict[str, str],
    seed: int,
) -> None:
    headers = {cell.value: cell.column for cell in sheet[DATA_HEADER_ROW]}
    key = hidden.set_index("Row_ID")
    for row_index in range(DATA_START_ROW, sheet.max_row + 1):
        row_id = str(sheet.cell(row_index, headers["Row_ID"]).value)
        item = key.loc[row_id]
        term = str(item["Term"])
        assigned = str(item["Assigned_Category"])
        expected = reference[_normalise(term)]
        if assigned == expected:
            base_level = 2
        elif tiers.get(assigned) == tiers.get(expected):
            base_level = 1
        else:
            base_level = 0
        if str(item["Final_Fate"]) in {"DEMOTED", "CRITIC_EXCLUDED", "CQ_EXCLUDED"}:
            base_level = min(base_level, 1)
        verdict = _verdict(base_level, expert_id, row_id, "category", seed)
        sheet.cell(row_index, headers["Correct (Yes/Partial/No/Unsure)"], verdict)
        sheet.cell(row_index, headers["Notes"], f"[{PROVENANCE}; A-anchored pseudo-oracle]")


def _fill_final_sheets(
    workbook,
    key: pd.DataFrame,
    expert_id: str,
    seed: int,
) -> None:
    def key_for(sheet_name: str) -> pd.DataFrame:
        return key[(key["Expert_ID"] == expert_id) & (key["Sheet"] == sheet_name)].set_index("Row_ID")

    sheet = workbook["Taxonomy"]
    headers = {cell.value: cell.column for cell in sheet[DATA_HEADER_ROW]}
    hidden = key_for("Taxonomy")
    for row_index in range(DATA_START_ROW, sheet.max_row + 1):
        row_id = str(sheet.cell(row_index, headers["Row_ID"]).value)
        item = hidden.loc[row_id]
        intermediate = str(item.get("Is_Intermediate", "False")).lower() == "true"
        relationship = _verdict(1 if intermediate else 2, expert_id, row_id, "taxonomy", seed)
        keep = _verdict(0 if intermediate else 2, expert_id, row_id, "taxonomy-core", seed, partial=False)
        sheet.cell(row_index, headers["Relationship_Correct (Yes/Partial/No/Unsure)"], relationship)
        sheet.cell(row_index, headers["Useful_PreSalt_Distinction (Yes/No/Unsure)"], keep)
        sheet.cell(row_index, headers["Notes"], f"[{PROVENANCE}]")

    sheet = workbook["Defined_Classes"]
    headers = {cell.value: cell.column for cell in sheet[DATA_HEADER_ROW]}
    hidden = key_for("Defined_Classes")
    for row_index in range(DATA_START_ROW, sheet.max_row + 1):
        row_id = str(sheet.cell(row_index, headers["Row_ID"]).value)
        item = hidden.loc[row_id]
        confidence_level = 2 if str(item.get("Definition_Type", "")) == "bearer_realizable" else 1
        definition = _verdict(confidence_level, expert_id, row_id, "definition", seed)
        definition = {
            "Yes": "Correct",
            "Partial": "Partly correct",
            "No": "Incorrect",
            "Unsure": "Unsure",
        }[definition]
        sheet.cell(
            row_index,
            headers["Definition_Verdict (Correct/Partly correct/Incorrect/Unsure)"],
            definition,
        )
        if definition == "Partly correct":
            sheet.cell(
                row_index,
                headers["Issue_Reason (select for Partly/Incorrect)"],
                "Feature is not defining",
            )
        elif definition == "Incorrect":
            sheet.cell(
                row_index,
                headers["Issue_Reason (select for Partly/Incorrect)"],
                "Base kind is wrong",
            )
        sheet.cell(row_index, headers["Notes"], f"[{PROVENANCE}]")

    sheet = workbook["Relations"]
    headers = {cell.value: cell.column for cell in sheet[DATA_HEADER_ROW]}
    hidden = key_for("Relations")
    for row_index in range(DATA_START_ROW, sheet.max_row + 1):
        row_id = str(sheet.cell(row_index, headers["Row_ID"]).value)
        item = hidden.loc[row_id]
        confidence = float(item.get("Confidence", 0.8))
        scope_confidence = float(item.get("Scope_Confidence", 0.8))
        needs_review = str(item.get("Scope_Needs_Review", "False")).lower() == "true"
        statement = _verdict(2 if confidence >= 0.8 else 1, expert_id, row_id, "relation", seed)
        scope = _verdict(
            2 if scope_confidence >= 0.85 and not needs_review else 1,
            expert_id,
            row_id,
            "relation-scope",
            seed,
            partial=False,
        )
        if "Unsure" in {statement, scope}:
            verdict = "Unsure"
        elif statement == "No":
            verdict = "Incorrect"
        elif statement == "Partial":
            verdict = "Partly wrong"
        elif scope == "No":
            verdict = "Context-specific"
        else:
            verdict = "Generally true"
        sheet.cell(row_index, headers["Relation_Verdict"], verdict)
        sheet.cell(row_index, headers["Notes"], f"[{PROVENANCE}; confidence-derived]")

    sheet = workbook["Individuals"]
    headers = {cell.value: cell.column for cell in sheet[DATA_HEADER_ROW]}
    hidden = key_for("Individuals")
    for row_index in range(DATA_START_ROW, sheet.max_row + 1):
        row_id = str(sheet.cell(row_index, headers["Row_ID"]).value)
        item = hidden.loc[row_id]
        named = _verdict(2, expert_id, row_id, "individual", seed, partial=False)
        type_level = 1 if pd.notna(item.get("Mint_Parent")) and str(item.get("Mint_Parent")).strip() else 2
        type_verdict = _verdict(type_level, expert_id, row_id, "individual-type", seed)
        sheet.cell(row_index, headers["Specific_Named_Entity (Yes/No/Unsure)"], named)
        sheet.cell(row_index, headers["Type_Correct (Yes/Partial/No/Unsure)"], type_verdict)
        sheet.cell(row_index, headers["Notes"], f"[{PROVENANCE}]")

    sheet = workbook["Meaning_Preservation"]
    headers = {cell.value: cell.column for cell in sheet[DATA_HEADER_ROW]}
    hidden = key_for("Meaning_Preservation")
    for row_index in range(DATA_START_ROW, sheet.max_row + 1):
        row_id = str(sheet.cell(row_index, headers["Row_ID"]).value)
        item = hidden.loc[row_id]
        confidence = float(item.get("confidence", 0.8))
        needs_review = str(item.get("needs_review", "False")).lower() == "true"
        level = 2 if confidence >= 0.9 and not needs_review else 1
        agree = _verdict(level, expert_id, row_id, "critic", seed)
        preservation = {
            "Yes": "Fully",
            "Partial": "Mostly",
            "No": "Reject",
            "Unsure": "Unsure",
        }[agree]
        if preservation == "Reject":
            preservation = "No"
        sheet.cell(
            row_index,
            headers["Meaning_Preserved (Fully/Mostly/No/Unsure)"],
            preservation,
        )
        decision_type = str(item.get("Decision_Type"))
        if preservation == "No":
            treatment = "Keep as separate concept"
        elif preservation == "Mostly":
            treatment = (
                "Leave out"
                if decision_type == "EXCLUDE"
                else "Keep information but not as separate concept"
            )
        else:
            treatment = ""
        sheet.cell(
            row_index,
            headers["Preferred_Outcome (for Mostly/No)"],
            treatment,
        )
        sheet.cell(row_index, headers["Notes"], f"[{PROVENANCE}; confidence-derived]")


def _fill_mock_workbooks(
    workbook_paths: list[str],
    key_path: str,
    reference: dict[str, str],
    seed: int,
) -> str:
    key = read_csv(key_path)
    key["Data_Provenance"] = PROVENANCE
    key["Mock_Ratings"] = True
    key["Mock_Seed"] = seed
    write_csv(key, key_path)
    _, tiers, _ = _category_metadata()

    for expert_number, workbook_path in enumerate(workbook_paths, 1):
        expert_id = f"expert_{expert_number}"
        workbook = load_workbook(workbook_path)
        _add_rehearsal_warning(workbook, expert_id)
        representation_key = key[
            (key["Expert_ID"] == expert_id) & (key["Sheet"] == "Representation")
        ]
        category_key = key[
            (key["Expert_ID"] == expert_id) & (key["Sheet"] == "Category_Correct")
        ]
        _fill_representation_sheet(
            workbook["Representation"], representation_key, expert_id, seed
        )
        _fill_category_sheet(
            workbook["Category_Correct"],
            category_key,
            expert_id,
            reference,
            tiers,
            seed,
        )
        _fill_final_sheets(workbook, key, expert_id, seed)
        timing = workbook["Timing"]
        timing_headers = {
            cell.value: cell.column for cell in timing[DATA_HEADER_ROW]
        }
        for row_index in range(DATA_START_ROW, timing.max_row + 1):
            row_id = str(timing.cell(row_index, timing_headers["Row_ID"]).value)
            minutes = 35 + int(_unit(seed, expert_id, row_id, "timing") * 31)
            timing.cell(row_index, timing_headers["Minutes"], minutes)
            timing.cell(row_index, timing_headers["Comments"], f"[{PROVENANCE}]")
        workbook.save(workbook_path)
    return key_path


def _verify_outputs(
    output_dir: Path,
    a_nld_path: str,
    a_category_path: str,
    workbook_paths: list[str],
    key_path: str,
) -> list[str]:
    checks = []
    if _sha256_file(a_nld_path) != _sha256_file(output_dir / "nld_A.csv"):
        raise AssertionError("Frozen A NLD copy changed")
    checks.append("Frozen A NLD copied byte-for-byte")
    if _sha256_file(a_category_path) != _sha256_file(output_dir / "cat_A.csv"):
        raise AssertionError("Frozen A category copy changed")
    checks.append("Frozen A categories copied byte-for-byte")

    term_sets = []
    for condition in CONDITIONS:
        nld = read_csv(output_dir / f"nld_{condition}.csv")
        category = read_csv(output_dir / f"cat_{condition}.csv")
        if len(nld) != EXPECTED_TERM_COUNT or len(category) != EXPECTED_TERM_COUNT:
            raise AssertionError(f"Condition {condition} row count mismatch")
        if nld["Term"].duplicated().any() or category["Term"].duplicated().any():
            raise AssertionError(f"Condition {condition} has duplicate terms")
        if category["Category"].astype(str).str.startswith("ERROR").any():
            raise AssertionError(f"Condition {condition} contains ERROR categories")
        term_sets.append(set(nld["Term"].astype(str)))
    if any(term_set != term_sets[0] for term_set in term_sets[1:]):
        raise AssertionError("Condition term sets differ")
    checks.append("All four conditions contain the same 407 unique terms")

    nld_a = read_csv(output_dir / "nld_A.csv").set_index("Term")
    nld_d = read_csv(output_dir / "nld_D.csv").set_index("Term")
    if not nld_a["Context"].equals(nld_d["Context"]):
        raise AssertionError("Condition D does not reuse exact A contexts")
    checks.append("Condition D reuses every frozen A context exactly")

    merged = read_csv(output_dir / "ablation_merged.csv")
    if len(merged) != EXPECTED_TERM_COUNT * 4:
        raise AssertionError("Merged matrix does not contain 1628 rows")
    checks.append("Merged Layer 1 matrix has 1628 unique term-condition rows")

    key = read_csv(key_path)
    if not key["Data_Provenance"].eq(PROVENANCE).all() or not key["Mock_Ratings"].eq(True).all():
        raise AssertionError("Blinding key lacks mock provenance")
    for path in workbook_paths:
        workbook = load_workbook(path, read_only=True, data_only=True)
        if workbook.sheetnames[0] != "REHEARSAL_ONLY":
            raise AssertionError(f"Workbook lacks leading warning sheet: {path}")
        forbidden_headers = {
            "Suggested_Category",
            "Suggested_Parent",
            "Suggested_Type",
            "Suggested_Change",
        }
        for sheet_name in (
            "Category_Correct",
            "Taxonomy",
            "Defined_Classes",
            "Individuals",
        ):
            headers = {cell.value for cell in next(workbook[sheet_name].iter_rows(max_row=1))}
            if headers & forbidden_headers:
                raise AssertionError(
                    f"{sheet_name} asks for a hidden-vocabulary replacement: "
                    f"{sorted(headers & forbidden_headers)}"
                )
        if Path(path).parent.name != "expert_workbooks":
            raise AssertionError(f"Distributable workbook is not isolated: {path}")
    if Path(key_path).parent.name != "private":
        raise AssertionError("Blinding key is not isolated in the private directory")
    checks.append("All workbooks and key are visibly marked as synthetic mock data")
    checks.append("No expert task requests a category, parent, type, or correction from an unseen list")
    checks.append("Distributable workbooks and private blinding key are separated")

    if not (output_dir / "analysis" / "layer2" / "layer2_results.json").exists():
        raise AssertionError("Layer 2 analysis output is missing")
    checks.append("Layer 1 and Layer 2 analyses completed without Azure calls")
    return checks


def _format_number(value: object, digits: int = 4) -> str:
    if value is None or (isinstance(value, float) and not np.isfinite(value)):
        return "n/a"
    return f"{float(value):.{digits}f}"


def _format_p(value: object) -> str:
    if value is None:
        return "n/a"
    number = float(value)
    return f"{number:.3e}" if number < 0.0001 else f"{number:.8f}"


def _write_findings_report(
    output_dir: Path,
    layer1_results: dict,
    layer2_results: dict,
    checks: list[str],
    categories: dict[str, pd.DataFrame],
    workbook_manifest: dict,
) -> Path:
    exact = layer1_results["exact_agreement"]
    tier = layer1_results["tier_agreement"]
    sensitivity = layer1_results["sensitivity_flags"]
    cochran = layer1_results["cochrans_q"].iloc[0]
    mcnemar = layer1_results["mcnemar_posthoc"]
    stuart_maxwell = layer1_results["stuart_maxwell"]
    category_result = layer2_results["category_correctness"]
    nld_result = layer2_results["representation"]["nld_quality"]
    final = layer2_results["final_ontology"]

    a_exact = exact[exact["Pair"].isin(["A vs B", "A vs C", "A vs D"])]
    a_tier = tier[tier["Pair"].isin(["A vs B", "A vs C", "A vs D"])]
    perturbation_counts = {
        condition: int(frame["Perturbed"].sum())
        for condition, frame in categories.items()
        if condition != "A"
    }

    workload_issues = []
    category_rows = int(workbook_manifest["category_rows"])
    if not 130 <= category_rows <= 200:
        workload_issues.append(
            f"Category module has {category_rows} rows, outside the planned 130-200 range."
        )
    else:
        workload_issues.append(
            f"Category module has {category_rows} rows, inside the planned 130-200 range."
        )

    final_outcomes = [
        ("Taxonomy relationship", final["taxonomy"]["relationship_correctness"]),
        ("Useful Pre-Salt distinction", final["taxonomy"]["useful_presalt_distinction"]),
        ("Defined-class verdict", final["defined_classes"]["definition_verdict"]),
        ("Relation verdict", final["relations"]["relation_verdict"]),
        ("Named entity", final["individuals"]["named_entity_correctness"]),
        ("Individual type", final["individuals"]["type_correctness"]),
    ]
    final_outcomes.extend(
        (f"Meaning preservation: {decision_type}", summary)
        for decision_type, summary in final["meaning_preservation"].items()
    )

    lines = [
        "# Offline Evaluation Rehearsal Findings",
        "",
        f"> **{WARNING}**",
        "",
        "## Scope",
        "",
        "This run validates the mechanics of the full evaluation workflow. B uses the first sentence of frozen A as a no-RAG proxy, B/C/D categories are deterministic perturbations of A, and workbook ratings are deterministic mock judgments. Therefore every p-value, confidence interval, effect size, and apparent condition advantage below is an expected property of the simulation, not empirical evidence.",
        "",
        "## Process Checks",
        "",
        *[f"- PASS: {check}" for check in checks],
        "",
        "## Synthetic Ablation Shape",
        "",
        f"- Perturbed assignments: B={perturbation_counts['B']}, C={perturbation_counts['C']}, D={perturbation_counts['D']} of 407.",
        f"- Independent sensitivities: RAG={int(sensitivity['RAG_Sensitive'].sum())}, NLD={int(sensitivity['NLD_Sensitive'].sum())}, structuring={int(sensitivity['Structuring_Sensitive'].sum())}.",
        f"- Cochran Q: Q({int(cochran['df'])})={_format_number(cochran['Q'])}, p={_format_p(cochran['p_value'])}.",
        "",
        "### Exact agreement",
        "",
        "| Pair | Agreement | Cohen kappa |",
        "|---|---:|---:|",
        *[
            f"| {row.Pair} | {_format_number(row.Agreement_Rate)} | {_format_number(row.Cohens_Kappa)} |"
            for row in a_exact.itertuples()
        ],
        "",
        "### Tier agreement",
        "",
        "| Pair | Agreement | Cohen kappa |",
        "|---|---:|---:|",
        *[
            f"| {row.Pair} | {_format_number(row.Agreement_Rate)} | {_format_number(row.Cohens_Kappa)} |"
            for row in a_tier.itertuples()
        ],
        "",
        "### Agreement-rate post-hoc contrasts",
        "",
        "| Comparison | Rate difference | Discordant pairs | Holm p | Significant |",
        "|---|---:|---:|---:|---|",
        *[
            f"| {row.Comparison} | {_format_number(row.Rate_Difference)} | {row.Discordant} | {_format_p(row.p_value_holm)} | {row.Significant} |"
            for row in mcnemar.itertuples()
        ],
        "",
        "### Tier marginal-homogeneity tests",
        "",
        "| Pair | Statistic | Effective df | Levels | Rank deficient | Holm p |",
        "|---|---:|---:|---:|---|---:|",
        *[
            f"| {row.Pair} | {_format_number(row.Statistic)} | {row.df} | {row.Levels} | {row.Rank_Deficient} | {_format_p(row.p_value_holm)} |"
            for row in stuart_maxwell.itertuples()
        ],
        "",
        "## Mock Expert Workload",
        "",
        f"- Representation terms: {workbook_manifest['representation_sample']}.",
        f"- Unique category assignments: {category_rows}.",
        *[f"- {name}: {count}." for name, count in workbook_manifest["final_sample_sizes"].items()],
        f"- Sample fates: {layer2_results['representation']['sample_final_fates']}.",
        "",
        "## Mock Layer 2 Results",
        "",
        f"- NLD A-B mean difference: {_format_number(nld_result['mean_difference_A_minus_B'])}; Wilcoxon W={_format_number(nld_result['wilcoxon']['W'])}, p={_format_p(nld_result['wilcoxon']['p_value'])}, rank-biserial={_format_number(nld_result['wilcoxon']['rank_biserial'])}.",
        f"- Preference: {nld_result['preference_sign_test']}.",
        f"- Category Friedman: chi2({category_result['friedman']['df']})={_format_number(category_result['friedman']['chi2'])}, p={_format_p(category_result['friedman']['p_value'])}, Kendall W={_format_number(category_result['friedman']['kendalls_w'])}; complete terms={category_result['friedman']['n_complete_terms']}, excluded={category_result['friedman']['n_excluded_incomplete_or_unsure']}.",
        f"- Category post-hoc diagnostics: {category_result['posthoc_wilcoxon_holm']}.",
        "",
        "### Mock category correctness",
        "",
        "| Condition | Mean score | Proportion yes | 95% CI |",
        "|---|---:|---:|---|",
        *[
            f"| {condition} | {_format_number(summary['mean_score'])} | {_format_number(summary['proportion_yes'])} | {summary['proportion_yes_ci_95']} |"
            for condition, summary in category_result["correctness_by_condition"].items()
        ],
        "",
        "### Mock final-ontology outcomes",
        "",
        "| Outcome | Items | Mean score | Positive proportion | Positive 95% CI | Unsure rate | Kappa | Prevalence warning |",
        "|---|---:|---:|---:|---|---:|---:|---|",
        *[
            f"| {name} | {summary['n_items']} | {_format_number(summary['mean_score'])} | {_format_number(summary['proportion_positive'])} | {summary['proportion_positive_ci_95']} | {_format_number(summary['unsure_rate'])} | {_format_number(summary['fleiss_kappa']['kappa'])} | {summary['fleiss_kappa']['prevalence_warning']} |"
            for name, summary in final_outcomes
        ],
        "",
        "## Issues and Improvements Identified",
        "",
        "1. **Actual B/C/D results remain unknowable offline.** The rehearsal confirms data flow only; no simulated statistic belongs in the thesis Results chapter.",
        "2. **The B proxy leaks Condition A.** Its first sentence comes from the RAG-grounded A definition, so it cannot estimate the causal RAG contribution. It is useful only for workbook and analyzer testing.",
        "3. **The mock category oracle is A-anchored.** Apparent expert support for A is built into the rehearsal. This verifies unblinding and statistical direction, not correctness.",
        "4. **Mock final-ontology judgments use critic confidence.** They cannot independently validate meaning preservation, relation scope, or retained-core quality.",
        "5. **Significance is easy to manufacture with 407 paired rows.** Real reporting must emphasize agreement rates, confusion patterns, and effect sizes rather than treating small p-values as accuracy evidence.",
        "6. **Context_Used remains descriptive.** The eight A rows with `Context_Used=False` are too few for a credible causal subgroup claim.",
        "7. **A real pilot is still required.** Have the three planned geology specialties complete a small subset before final-study distribution; check wording, fatigue, completion time, and use of Partial versus Unsure.",
        "8. **Sparse post-hoc contrasts need explicit review.** A pairwise result can be significant when only a handful of terms have nonzero score differences; report `n_nonzero_pairs`, its fraction, and the mean difference together.",
        "9. **Kappa is prevalence-sensitive.** Report rating marginals and raw agreement alongside kappa, especially when Yes or Partial dominates.",
        "10. **Stuart-Maxwell may have reduced effective degrees of freedom.** This is valid when a tier has no discordant transitions; report the covariance rank and rank-deficiency flag rather than forcing `levels-1` degrees of freedom.",
        "11. **The private key is operationally sensitive.** Distribute only files from `expert_workbooks/`; keep `private/` inaccessible to experts.",
        *[f"12. **Workload check:** {issue}" for issue in workload_issues],
        "",
        "## Go/No-Go Before Paid Execution",
        "",
        "- GO: schemas, manifests, term pairing, workbook generation, stable row IDs, category propagation, mock completion, unblinding, bootstrap CIs, omnibus gates, and separate final-task analyses all execute offline.",
        "- NO-GO for interpretation: do not quote any rehearsal result as evidence for RAG, NLDs, structuring, or ontology quality.",
        "- Before Azure: archive or remove this rehearsal directory, run paid B/C/D in `evaluation_study/output/ablation/`, inspect ten random rows per condition, and verify the production manifest before analysis.",
    ]
    report_path = output_dir / "OFFLINE_REHEARSAL_FINDINGS.md"
    report_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return report_path


def run_offline_rehearsal(
    output_dir: str = DEFAULT_OUTPUT_DIR,
    ontology_dir: str = DEFAULT_ONTOLOGY_DIR,
    terms_path: str = DEFAULT_TERMS_PATH,
    a_nld_path: str = DEFAULT_A_NLD_PATH,
    a_category_path: str = DEFAULT_A_CATEGORY_PATH,
    seed: int = SEED,
    n_experts: int = 3,
    overwrite: bool = False,
    bootstrap_iterations: int = 5000,
) -> dict:
    """Execute the complete deterministic offline rehearsal."""
    started = datetime.now(timezone.utc)
    rehearsal_dir = _prepare_output_dir(output_dir, overwrite)
    terms, nld_a, cat_a = _load_source_inputs(terms_path, a_nld_path, a_category_path)
    nlds = _build_nld_surrogates(rehearsal_dir, nld_a, a_nld_path)
    categories, fates, reference = _build_category_surrogates(
        rehearsal_dir,
        terms,
        nlds,
        cat_a,
        a_category_path,
        ontology_dir,
        seed,
    )
    _write_merged(rehearsal_dir, categories)

    warning_path = rehearsal_dir / "DO_NOT_USE_AS_THESIS_RESULTS.md"
    warning_path.write_text(
        f"# {PROVENANCE}\n\n{WARNING}\n\n"
        "This directory exists to find process defects before paid execution. "
        "Delete or archive it before producing real study outputs.\n",
        encoding="utf-8",
    )

    layer1_results = _run_layer1(rehearsal_dir)
    workbook_paths, key_path = generate_expert_evaluation(
        n_terms=100,
        seed=seed,
        output_dir=str(rehearsal_dir),
        n_experts=n_experts,
        ontology_dir=ontology_dir,
        terms_path=terms_path,
        require_reviewed_definitions=False,
    )
    _fill_mock_workbooks(workbook_paths, key_path, reference, seed)
    layer2_dir = rehearsal_dir / "analysis" / "layer2"
    layer2_results = run_modular_analysis(
        workbook_paths,
        key_path,
        str(layer2_dir),
        bootstrap_iterations=bootstrap_iterations,
        seed=seed,
    )
    checks = _verify_outputs(
        rehearsal_dir,
        a_nld_path,
        a_category_path,
        workbook_paths,
        key_path,
    )

    workbook_manifest_path = rehearsal_dir / "expert_evaluation_manifest.json"
    workbook_manifest = json.loads(workbook_manifest_path.read_text(encoding="utf-8"))
    workbook_manifest["data_provenance"] = PROVENANCE
    workbook_manifest["mock_ratings"] = True
    workbook_manifest["warning"] = WARNING
    workbook_manifest_path.write_text(
        json.dumps(workbook_manifest, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    report_path = _write_findings_report(
        rehearsal_dir,
        layer1_results,
        layer2_results,
        checks,
        categories,
        workbook_manifest,
    )

    generated_files = sorted(
        str(path.relative_to(rehearsal_dir))
        for path in rehearsal_dir.rglob("*")
        if path.is_file()
    )
    generated_file_sha256 = {
        relative_path: _sha256_file(rehearsal_dir / relative_path)
        for relative_path in generated_files
    }
    completed = datetime.now(timezone.utc)
    manifest = {
        "artifact_kind": PROVENANCE,
        "warning": WARNING,
        "status": "complete",
        "azure_calls": 0,
        "human_experts": 0,
        "scientific_inference_permitted": False,
        "started_utc": started.isoformat(),
        "completed_utc": completed.isoformat(),
        "seed": seed,
        "mock_expert_count": n_experts,
        "bootstrap_iterations": bootstrap_iterations,
        "term_count": len(terms),
        "term_set_sha256": _sha256_text(
            "\n".join(sorted(terms["Readable_Term"].astype(str)))
        ),
        "source_sha256": {
            "terms": _sha256_file(terms_path),
            "frozen_a_nld": _sha256_file(a_nld_path),
            "frozen_a_category": _sha256_file(a_category_path),
            "ontology_config": _sha256_file(get_config()._source_path),
        },
        "surrogate_design": {
            "B_NLD": "first sentence of frozen A; intentionally A-leaking process proxy",
            "C_NLD": "empty string",
            "D_NLD": "exact frozen A Context",
            "B_C_D_categories": "deterministic A perturbations weighted by final fate",
            "mock_experts": "deterministic artifact-derived ratings with bounded rater variation",
            "change_rates": CHANGE_RATES,
            "fate_multipliers": FATE_MULTIPLIERS,
        },
        "fate_counts": pd.Series(fates).value_counts().sort_index().to_dict(),
        "checks": checks,
        "findings_report": str(report_path.relative_to(rehearsal_dir)),
        "generated_files": generated_files,
        "generated_file_sha256": generated_file_sha256,
    }
    manifest_path = rehearsal_dir / "OFFLINE_REHEARSAL_MANIFEST.json"
    manifest_path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2, allow_nan=False),
        encoding="utf-8",
    )
    print(f"\nOffline rehearsal complete: {rehearsal_dir}")
    print(f"  Azure calls: 0")
    print(f"  Human experts: 0")
    print(f"  Findings: {report_path}")
    print(f"  Manifest: {manifest_path}")
    return {
        "output_dir": str(rehearsal_dir),
        "manifest": str(manifest_path),
        "findings": str(report_path),
        "checks": checks,
        "layer1": layer1_results,
        "layer2": layer2_results,
    }


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Run the strictly offline synthetic evaluation rehearsal"
    )
    parser.add_argument("--output-dir", default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--ontology-dir", default=DEFAULT_ONTOLOGY_DIR)
    parser.add_argument("--terms", default=DEFAULT_TERMS_PATH)
    parser.add_argument("--a-nld", default=DEFAULT_A_NLD_PATH)
    parser.add_argument("--a-categories", default=DEFAULT_A_CATEGORY_PATH)
    parser.add_argument("--seed", type=int, default=SEED)
    parser.add_argument("--experts", type=int, default=3)
    parser.add_argument("--bootstrap-iterations", type=int, default=5000)
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args()
    run_offline_rehearsal(
        output_dir=args.output_dir,
        ontology_dir=args.ontology_dir,
        terms_path=args.terms,
        a_nld_path=args.a_nld,
        a_category_path=args.a_categories,
        seed=args.seed,
        n_experts=args.experts,
        overwrite=args.overwrite,
        bootstrap_iterations=args.bootstrap_iterations,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())